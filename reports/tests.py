import io

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse

from classes.models import SchoolClass, Subject
from sessions.models import Session
from trainers.models import TrainerProfile


def create_trainer(username, full_name, employee_id):
    user = User.objects.create_user(username=username, password="Trainer@123")
    return TrainerProfile.objects.create(
        user=user, employee_id=employee_id, full_name=full_name
    )


def create_session(trainer, school_class, subject, date, topic, **kwargs):
    return Session.objects.create(
        trainer=trainer,
        school_class=school_class,
        subject=subject,
        date=date,
        topic_taught=topic,
        **kwargs,
    )


def docx_text(content):
    from docx import Document

    doc = Document(io.BytesIO(content))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def docx_cell(content, header_label):
    """Value from a two-row table whose header row contains ``header_label``."""
    from docx import Document

    doc = Document(io.BytesIO(content))
    for table in doc.tables:
        header = [cell.text.strip() for cell in table.rows[0].cells]
        if header_label in header:
            column = header.index(header_label)
            return table.rows[1].cells[column].text.strip()
    return None


class TrainerReportDownloadTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.trainer = create_trainer("alice", "Alice Sharma", "E1")
        cls.other = create_trainer("bob", "Bob Verma", "E2")
        cls.admin = User.objects.create_superuser(username="boss", password="Admin@123")

        cls.klass = SchoolClass.objects.create(name="Class 8", section="A")
        cls.subject = Subject.objects.create(name="Mathematics")

        create_session(
            cls.trainer, cls.klass, cls.subject, "2026-01-05",
            "ALICE-SECRET-TOPIC", location="Room 1", students_present=12,
        )
        create_session(
            cls.trainer, cls.klass, cls.subject, "2026-02-10",
            "Alice second topic", location="Room 2", students_present=11,
        )
        create_session(
            cls.other, cls.klass, cls.subject, "2026-01-07",
            "BOB-SECRET-TOPIC", location="Room 9", students_present=20,
        )

    def _download(self, fmt="docx", **params):
        query = {"format": fmt}
        query.update(params)
        return self.client.get(reverse("reports:download"), query)

    def test_login_required(self):
        response = self._download()
        self.assertEqual(response.status_code, 302)
        self.assertIn("/accounts/login/", response.url)

    def test_trainer_can_download_pdf(self):
        self.client.login(username="alice", password="Trainer@123")
        response = self._download(fmt="pdf")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertIn("attachment", response["Content-Disposition"])
        self.assertTrue(response.content.startswith(b"%PDF"))
        self.assertGreater(len(response.content), 1000)

    def test_trainer_can_download_docx(self):
        self.client.login(username="alice", password="Trainer@123")
        response = self._download(fmt="docx")
        self.assertEqual(response.status_code, 200)
        self.assertIn(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            response["Content-Type"],
        )
        self.assertIn(".docx", response["Content-Disposition"])
        text = docx_text(response.content)
        self.assertIn("Alice Sharma", text)
        self.assertIn("ALICE-SECRET-TOPIC", text)

    def test_trainer_report_excludes_other_trainers_data(self):
        self.client.login(username="alice", password="Trainer@123")
        response = self._download(fmt="docx")
        text = docx_text(response.content)
        self.assertNotIn("BOB-SECRET-TOPIC", text)
        self.assertNotIn("Bob Verma", text)
        # Only Alice's sessions are counted.
        self.assertEqual(docx_cell(response.content, "Total Sessions"), "2")
        self.assertEqual(docx_cell(response.content, "Sessions"), "2")

    def test_trainer_param_cannot_widen_scope(self):
        self.client.login(username="alice", password="Trainer@123")
        text = docx_text(
            self._download(fmt="docx", trainer=self.other.pk).content
        )
        self.assertNotIn("BOB-SECRET-TOPIC", text)
        self.assertNotIn("Bob Verma", text)

    def test_date_filters_are_applied(self):
        self.client.login(username="alice", password="Trainer@123")
        response = self._download(fmt="docx", **{"from": "2026-02-01"})
        text = docx_text(response.content)
        self.assertIn("Alice second topic", text)
        self.assertNotIn("ALICE-SECRET-TOPIC", text)
        self.assertEqual(docx_cell(response.content, "Total Sessions"), "1")

    def test_class_and_subject_filters_are_applied(self):
        other_class = SchoolClass.objects.create(name="Class 9", section="B")
        other_subject = Subject.objects.create(name="Science")
        create_session(
            self.trainer, other_class, other_subject, "2026-03-01",
            "SCIENCE-TOPIC",
        )
        self.client.login(username="alice", password="Trainer@123")
        text = docx_text(
            self._download(
                fmt="docx",
                **{"class": other_class.pk, "subject": other_subject.pk},
            ).content
        )
        self.assertIn("SCIENCE-TOPIC", text)
        self.assertNotIn("ALICE-SECRET-TOPIC", text)

    def test_report_page_exposes_download_links(self):
        self.client.login(username="alice", password="Trainer@123")
        response = self.client.get(reverse("reports:report"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "format=pdf")
        self.assertContains(response, "format=docx")

    def test_dashboard_exposes_download_links(self):
        self.client.login(username="alice", password="Trainer@123")
        response = self.client.get(reverse("dashboard"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, reverse("reports:download"))
        self.assertNotContains(response, "BOB-SECRET-TOPIC")

    def test_user_without_profile_is_denied(self):
        User.objects.create_user(username="noprofile", password="Trainer@123")
        self.client.login(username="noprofile", password="Trainer@123")
        response = self._download()
        self.assertEqual(response.status_code, 403)


class AdminReportDownloadTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_superuser(username="boss", password="Admin@123")
        cls.trainer = create_trainer("alice", "Alice Sharma", "E1")
        cls.other = create_trainer("bob", "Bob Verma", "E2")
        cls.klass = SchoolClass.objects.create(name="Class 8", section="A")
        cls.subject = Subject.objects.create(name="Mathematics")
        create_session(
            cls.trainer, cls.klass, cls.subject, "2026-01-05", "ALICE-TOPIC"
        )
        create_session(
            cls.other, cls.klass, cls.subject, "2026-01-07", "BOB-TOPIC"
        )

    def test_admin_pdf_download_still_works(self):
        self.client.login(username="boss", password="Admin@123")
        response = self.client.get(
            reverse("reports:download"), {"format": "pdf"}
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF"))

    def test_admin_docx_contains_all_trainers(self):
        self.client.login(username="boss", password="Admin@123")
        response = self.client.get(
            reverse("reports:download"), {"format": "docx"}
        )
        self.assertEqual(response.status_code, 200)
        text = docx_text(response.content)
        self.assertIn("ALICE-TOPIC", text)
        self.assertIn("BOB-TOPIC", text)
        self.assertIn("Active Trainers", text)
        self.assertIn("Alice Sharma", text)
        self.assertIn("Bob Verma", text)

    def test_admin_trainer_filter_is_applied(self):
        self.client.login(username="boss", password="Admin@123")
        response = self.client.get(
            reverse("reports:download"),
            {"format": "docx", "trainer": self.other.pk},
        )
        text = docx_text(response.content)
        self.assertIn("BOB-TOPIC", text)
        self.assertNotIn("ALICE-TOPIC", text)

    def test_trainer_cannot_use_admin_report(self):
        self.client.login(username="alice", password="Trainer@123")
        response = self.client.get(reverse("reports:report"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "My Report")
        self.assertNotContains(response, "Active Trainers")
